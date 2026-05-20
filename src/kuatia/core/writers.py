"""Escritores de output: `.txt`, `.srt`, `.vtt`, `.docx`.

Recebem `Iterable[Segment]` para serem reutilizáveis por CLI e GUI sem
acoplamento ao pipeline de transcrição.
"""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document

from kuatia.core.transcriber import Segment


@dataclass(frozen=True)
class DocxMeta:
    """Metadados exibidos no cabeçalho do `.docx`.

    `generated_at` é injetável pra testabilidade — em produção, callers
    passam `datetime.now()`; em testes, passam um instante fixo.
    """

    input_name: str = ""
    duration_sec: float | None = None
    model_name: str = ""
    generated_at: datetime | None = None


def _format_timestamp(seconds: float) -> str:
    """Formato SRT: HH:MM:SS,mmm (vírgula nos milissegundos)."""
    seconds = max(0.0, float(seconds))
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = seconds - (h * 3600 + m * 60)
    return f"{h:02d}:{m:02d}:{s:06.3f}".replace(".", ",")


def _format_timestamp_vtt(seconds: float) -> str:
    """Formato WebVTT: HH:MM:SS.mmm (ponto nos milissegundos)."""
    return _format_timestamp(seconds).replace(",", ".")


def _format_duration(seconds: float) -> str:
    """Duração humana: `1h 23m 45s`, `23m 45s`, ou `45s`."""
    seconds = max(0.0, float(seconds))
    total = int(round(seconds))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}h {m:02d}m {s:02d}s"
    if m:
        return f"{m}m {s:02d}s"
    return f"{s}s"


def _format_datetime_br(dt: datetime) -> str:
    """Formato BR: `DD/MM/YYYY HH:MM`."""
    return dt.strftime("%d/%m/%Y %H:%M")


def write_txt(segments: Iterable[Segment], output_path: Path) -> None:
    """Uma linha por segment, com timestamps `[start --> end] texto`."""
    lines = [
        f"[{_format_timestamp(s.start)} --> {_format_timestamp(s.end)}] {s.text}" for s in segments
    ]
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_srt(segments: Iterable[Segment], output_path: Path) -> None:
    """SubRip — formato `.srt` padrão."""
    lines: list[str] = []
    for idx, s in enumerate(segments, start=1):
        lines.append(str(idx))
        lines.append(f"{_format_timestamp(s.start)} --> {_format_timestamp(s.end)}")
        lines.append(s.text)
        lines.append("")
    output_path.write_text("\n".join(lines), encoding="utf-8")


def write_vtt(segments: Iterable[Segment], output_path: Path) -> None:
    """WebVTT — formato `.vtt` para players HTML5."""
    lines: list[str] = ["WEBVTT", ""]
    for s in segments:
        lines.append(f"{_format_timestamp_vtt(s.start)} --> {_format_timestamp_vtt(s.end)}")
        lines.append(s.text)
        lines.append("")
    output_path.write_text("\n".join(lines), encoding="utf-8")


def write_docx(
    segments: Iterable[Segment],
    output_path: Path,
    meta: DocxMeta | None = None,
) -> None:
    """Documento Word com cabeçalho de metadados + transcrição em parágrafos.

    Layout: cabeçalho com info do input/modelo/duração/data seguido de
    parágrafos no formato `**[HH:MM:SS,mmm]** — texto`. Layout em parágrafos
    (vs tabela) foi escolhido por copiar limpo entre versões do Word e por
    ser mais natural pra continuar editando como nota.
    """
    meta = meta or DocxMeta()
    generated_at = meta.generated_at or datetime.now()

    doc = Document()
    doc.add_heading("Transcrição", level=1)

    info_lines = [
        ("Arquivo: ", meta.input_name or "—"),
        (
            "Duração: ",
            _format_duration(meta.duration_sec) if meta.duration_sec is not None else "—",
        ),
        ("Modelo: ", meta.model_name or "—"),
        ("Gerado em: ", _format_datetime_br(generated_at)),
    ]
    for label, value in info_lines:
        p = doc.add_paragraph()
        run = p.add_run(label)
        run.bold = True
        p.add_run(value)

    doc.add_paragraph()
    doc.add_heading("Conteúdo", level=2)
    for s in segments:
        p = doc.add_paragraph()
        ts_run = p.add_run(f"[{_format_timestamp(s.start)}] ")
        ts_run.bold = True
        p.add_run(f"— {s.text}")

    doc.save(str(output_path))
