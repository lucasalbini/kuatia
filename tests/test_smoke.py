"""Smoke tests do package — verifica importação e helpers puros sem dependência de modelo."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest


def test_package_importa() -> None:
    import kuatia

    assert kuatia.__doc__


def test_main_entrypoints_existem() -> None:
    from kuatia.cli import main as transcribe_main
    from kuatia.convert_model import main as convert_main

    assert callable(convert_main)
    assert callable(transcribe_main)


def test_core_api_importa() -> None:
    from kuatia.core.audio import SAMPLE_RATE, load_audio
    from kuatia.core.errors import (
        AudioLoadError,
        KuatiaError,
        ModelNotFoundError,
        TranscriptionError,
    )
    from kuatia.core.transcriber import Segment, Transcriber
    from kuatia.core.writers import DocxMeta, write_docx, write_srt, write_txt, write_vtt

    assert SAMPLE_RATE == 16_000
    assert callable(load_audio)
    assert callable(Transcriber)
    assert callable(write_txt)
    assert callable(write_srt)
    assert callable(write_vtt)
    assert callable(write_docx)
    assert callable(DocxMeta)
    assert issubclass(AudioLoadError, KuatiaError)
    assert issubclass(ModelNotFoundError, KuatiaError)
    assert issubclass(TranscriptionError, KuatiaError)
    assert Segment(0.0, 1.0, "x").text == "x"


def test_write_docx_estrutura(tmp_path: Path) -> None:
    from docx import Document

    from kuatia.core.transcriber import Segment
    from kuatia.core.writers import DocxMeta, write_docx

    segments = [
        Segment(0.0, 1.5, "Olá mundo"),
        Segment(2.0, 3.5, "Segunda fala"),
        Segment(4.0, 5.5, "Terceira fala"),
    ]
    meta = DocxMeta(
        input_name="exemplo.mp4",
        duration_sec=125.5,
        model_name="whisper-large-v3",
        generated_at=datetime(2026, 5, 20, 22, 45),
    )
    out = tmp_path / "out.docx"
    write_docx(segments, out, meta=meta)

    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # Cabeçalho deve conter os campos do meta
    assert "Transcrição" in full_text
    assert "exemplo.mp4" in full_text
    assert "whisper-large-v3" in full_text
    assert "20/05/2026" in full_text
    assert "22:45" in full_text
    assert "2m 06s" in full_text  # 125.5s arredonda pra 126s = 2m 06s

    # Conteúdo: 3 segments aparecem no documento
    assert "Olá mundo" in full_text
    assert "Segunda fala" in full_text
    assert "Terceira fala" in full_text

    # Timestamps em formato SRT
    assert "[00:00:00,000]" in full_text
    assert "[00:00:02,000]" in full_text
    assert "[00:00:04,000]" in full_text


def test_write_docx_sem_meta(tmp_path: Path) -> None:
    """Sem `meta`, doc usa placeholders e datetime.now() pra `Gerado em`."""
    from docx import Document

    from kuatia.core.transcriber import Segment
    from kuatia.core.writers import write_docx

    out = tmp_path / "minimal.docx"
    write_docx([Segment(0.0, 1.0, "fala única")], out)

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Transcrição" in full_text
    assert "fala única" in full_text
    assert "—" in full_text  # placeholder dos campos sem meta


@pytest.mark.parametrize(
    ("seconds", "expected"),
    [
        (0, "00:00:00,000"),
        (0.5, "00:00:00,500"),
        (61.5, "00:01:01,500"),
        (3661.123, "01:01:01,123"),
        (-1, "00:00:00,000"),  # valores negativos viram 0
    ],
)
def test_format_timestamp(seconds: float, expected: str) -> None:
    from kuatia.core.writers import _format_timestamp

    assert _format_timestamp(seconds) == expected


def test_chunks_to_segments_filtra_vazios_e_sem_start() -> None:
    from kuatia.core.transcriber import Segment, _chunks_to_segments

    chunks = [
        {"timestamp": (0.0, 1.0), "text": "olá"},
        {"timestamp": (None, 2.0), "text": "ignorado (sem start)"},
        {"timestamp": (3.0, None), "text": "fim faltando"},
        {"timestamp": (4.0, 5.0), "text": "   "},  # texto vazio após strip
        {"timestamp": (6.0, 7.0), "text": "mundo"},
    ]
    segments = _chunks_to_segments(chunks)
    assert segments == [
        Segment(0.0, 1.0, "olá"),
        Segment(3.0, 4.0, "fim faltando"),
        Segment(6.0, 7.0, "mundo"),
    ]
