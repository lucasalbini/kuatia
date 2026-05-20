"""Tests dos escritores de output (.txt, .srt, .vtt, .docx) e helpers de formatação."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from docx import Document

from kuatia.core.transcriber import Segment
from kuatia.core.writers import (
    DocxMeta,
    _format_datetime_br,
    _format_duration,
    _format_timestamp,
    _format_timestamp_vtt,
    write_docx,
    write_srt,
    write_txt,
    write_vtt,
)


@pytest.mark.parametrize(
    ("seconds", "expected"),
    [
        (0, "00:00:00,000"),
        (0.5, "00:00:00,500"),
        (61.5, "00:01:01,500"),
        (3661.123, "01:01:01,123"),
        (-1, "00:00:00,000"),
    ],
)
def test_format_timestamp(seconds: float, expected: str) -> None:
    assert _format_timestamp(seconds) == expected


@pytest.mark.parametrize(
    ("seconds", "expected"),
    [
        (0, "00:00:00.000"),
        (61.5, "00:01:01.500"),
        (3661.123, "01:01:01.123"),
    ],
)
def test_format_timestamp_vtt(seconds: float, expected: str) -> None:
    assert _format_timestamp_vtt(seconds) == expected


@pytest.mark.parametrize(
    ("seconds", "expected"),
    [
        (0, "0s"),
        (5, "5s"),
        (59.4, "59s"),
        (60, "1m 00s"),
        (65, "1m 05s"),
        (3600, "1h 00m 00s"),
        (3725, "1h 02m 05s"),
        (-10, "0s"),
    ],
)
def test_format_duration(seconds: float, expected: str) -> None:
    assert _format_duration(seconds) == expected


def test_format_datetime_br() -> None:
    dt = datetime(2026, 5, 20, 22, 45)
    assert _format_datetime_br(dt) == "20/05/2026 22:45"


def test_write_txt(tmp_path: Path, sample_segments: list[Segment]) -> None:
    out = tmp_path / "out.txt"
    write_txt(sample_segments, out)
    content = out.read_text(encoding="utf-8")
    assert "[00:00:00,000 --> 00:00:01,500] Olá mundo" in content
    assert "[00:00:02,000 --> 00:00:03,500] Segunda fala com acentos: ção, ã" in content
    assert "[00:00:04,000 --> 00:00:05,500] Terceira fala" in content
    assert content.endswith("\n")


def test_write_txt_vazio(tmp_path: Path) -> None:
    out = tmp_path / "vazio.txt"
    write_txt([], out)
    assert out.read_text(encoding="utf-8") == "\n"


def test_write_srt(tmp_path: Path, sample_segments: list[Segment]) -> None:
    out = tmp_path / "out.srt"
    write_srt(sample_segments, out)
    content = out.read_text(encoding="utf-8")
    # Cada bloco SRT: índice / timestamp / texto / linha em branco
    assert content.startswith("1\n00:00:00,000 --> 00:00:01,500\nOlá mundo\n")
    assert "2\n00:00:02,000 --> 00:00:03,500\nSegunda fala com acentos: ção, ã\n" in content
    assert "3\n00:00:04,000 --> 00:00:05,500\nTerceira fala\n" in content


def test_write_vtt(tmp_path: Path, sample_segments: list[Segment]) -> None:
    out = tmp_path / "out.vtt"
    write_vtt(sample_segments, out)
    content = out.read_text(encoding="utf-8")
    assert content.startswith("WEBVTT\n\n")
    # VTT usa ponto, não vírgula, nos ms
    assert "00:00:00.000 --> 00:00:01.500\nOlá mundo" in content
    assert "00:00:04.000 --> 00:00:05.500\nTerceira fala" in content
    # Garante que não tem vírgula nos timestamps
    assert ",500" not in content


def test_write_docx_estrutura(
    tmp_path: Path, sample_segments: list[Segment], fixed_meta: DocxMeta
) -> None:
    out = tmp_path / "out.docx"
    write_docx(sample_segments, out, meta=fixed_meta)

    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # Cabeçalho
    assert "Transcrição" in full_text
    assert "exemplo.mp4" in full_text
    assert "whisper-large-v3" in full_text
    assert "20/05/2026 22:45" in full_text
    assert "2m 06s" in full_text  # 125.5s → 126s → 2m 06s

    # Conteúdo
    assert "Conteúdo" in full_text
    assert "Olá mundo" in full_text
    assert "Segunda fala com acentos: ção, ã" in full_text
    assert "Terceira fala" in full_text

    # Timestamps no formato esperado
    assert "[00:00:00,000]" in full_text
    assert "[00:00:02,000]" in full_text
    assert "[00:00:04,000]" in full_text


def test_write_docx_sem_meta(tmp_path: Path) -> None:
    """Sem `meta`, doc usa placeholder `—` nos campos e `datetime.now()`."""
    out = tmp_path / "minimal.docx"
    write_docx([Segment(0.0, 1.0, "fala única")], out)

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Transcrição" in full_text
    assert "fala única" in full_text
    assert "—" in full_text


def test_write_docx_duration_none(tmp_path: Path) -> None:
    """`duration_sec=None` deve render placeholder, não crashar."""
    meta = DocxMeta(input_name="x.mp3", duration_sec=None, model_name="m")
    out = tmp_path / "no_dur.docx"
    write_docx([Segment(0.0, 1.0, "a")], out, meta=meta)

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Duração: —" in full_text


def test_write_docx_timestamp_bold(
    tmp_path: Path, sample_segments: list[Segment], fixed_meta: DocxMeta
) -> None:
    """Garante que o run do timestamp está em bold (sustenta a decisão de layout)."""
    out = tmp_path / "bold.docx"
    write_docx(sample_segments, out, meta=fixed_meta)

    doc = Document(str(out))
    # Procura pelo primeiro parágrafo de conteúdo (após cabeçalho + heading "Conteúdo")
    content_paragraphs = [p for p in doc.paragraphs if p.text.startswith("[")]
    assert content_paragraphs, "esperava ao menos um parágrafo de conteúdo"
    first = content_paragraphs[0]
    bold_runs = [r for r in first.runs if r.bold]
    assert bold_runs, "timestamp do primeiro segment deveria estar em bold"
    assert "[00:00:00,000]" in bold_runs[0].text
